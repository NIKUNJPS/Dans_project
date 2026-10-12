"""Export builders on a large, realistic MTO report.

A 300-row, 16-column take-off with characters that break naive builders ('<', '&', '"',
×, ≈, Σ) must produce every format quickly — the Word export used to take minutes and
the PDF crashed on '<', which made every download time out.
"""
import json
import os
import random
import tempfile
import time
import zipfile
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "test")
os.environ.setdefault("UPLOAD_DIR", tempfile.mkdtemp(prefix="sm_uploads_"))

import pytest  # noqa: E402

import export_service as E  # noqa: E402
from estimation import takeoff as T  # noqa: E402

PROFILES = ["W12x26", "W14x30", "W18x50", "W24x76", "HSS6x6x3/8", "HSS8x4x1/4", "C10x15.3",
            "L4x4x3/8", "L6x4x1/2", "WT6x13", "PIPE4STD", "PL1/2x8", "W8x31", "W10x33"]
FLAGS = ["", "", "LENGTH <SCALED> & VERIFY", 'Σ ≈ 2× per GL-3 — confirm ≥ 6"', "TBD"]


def _big_mto(batches: int = 12, rows_per_batch: int = 26) -> str:
    rnd = random.Random(7)
    parsed = []
    for b in range(batches):
        rows = ["no|type|role|mark|profile|qty|raw_length|length_mm|unit_wt_kgm|weight_kg|grade|finish|sheet|detail|method|flag"]
        for i in range(rows_per_batch):
            p = rnd.choice(PROFILES)
            rows.append(f"{i + 1}|MISC|PRIMARY|B{b}{i:02d}|{p}|{rnd.randint(1, 12)}|{rnd.randint(4, 40)}'-{rnd.randint(0, 11)}\"|"
                        f"|||A992|PRIMER & GALV|S-{b + 2}{i % 9:02d}|{i % 7}/S-50{i % 5}|SCHEDULE|{rnd.choice(FLAGS)}")
        meta = {"jurisdiction": "USA",
                "sheets": [{"sheet": f"S-{b + 2}{s:02d}", "title": f"Framing Plan {s} — Area <B> & C"} for s in range(9)],
                "rfis": [{"id": f"RFI-{k}", "priority": "HIGH", "sheet": f"S-{b + 2}01",
                          "question": f"Confirm B{b}{k:02d} at grid <A-3> & EL +12'-6\" (W12×26 vs W14×30)."}
                         for k in range(14)]}
        text = f"<<<META>>>\n{json.dumps(meta)}\n<<<END_META>>>\n<<<ROWS>>>\n" + "\n".join(rows) + "\n<<<END_ROWS>>>"
        parsed.append({"meta": T.parse_meta(text), "rows": T.parse_rows(text), "batch": b + 1})
    docs = [{"index": i + 1, "name": f"Binder_{i + 1}.pdf", "pages": 40, "size_mb": 32.3} for i in range(5)]
    return T.render_mto_report(T.consolidate(parsed, jurisdiction_hint="USA", documents=docs))


@pytest.fixture()
def export_dir(monkeypatch, tmp_path):
    monkeypatch.setattr(E, "EXPORT_DIR", tmp_path)
    return tmp_path


META = {"id": "rep1", "mode_label": "Master MTO Engine", "project_name": "Warehouse <North> & Annex",
        "completed_at": "2026-10-05", "model_used": "STRUCTMIND CORE · PRO", "blockchain_hash": "abc"}


@pytest.mark.parametrize("fmt", list(E.EXPORTERS))
def test_every_format_builds_fast_on_a_large_takeoff(export_dir, fmt):
    md = _big_mto()
    assert md.count("\n|") > 300
    t0 = time.perf_counter()
    path = Path(E.generate_export(fmt, md, META))
    elapsed = time.perf_counter() - t0
    assert path.exists() and path.stat().st_size > 1000
    assert elapsed < 30, f"{fmt} took {elapsed:.1f}s"
    assert not list(export_dir.glob(".*.tmp")), "temporary files left behind"


def test_pdf_and_word_keep_special_characters(export_dir):
    md = "## RFIs\nTable: Checks\n| Item | Note |\n|---|---|\n| B1 <A-3> | PRIMER & GALV, L < 6 m, 2× W12×26 |\n"
    pdf = Path(E.export_pdf(md, META))
    assert pdf.read_bytes()[:4] == b"%PDF"
    docx = Path(E.export_docx(md, META))
    with zipfile.ZipFile(docx) as z:
        xml = z.read("word/document.xml").decode()
    assert "B1 &lt;A-3&gt;" in xml and "PRIMER &amp; GALV" in xml and "W12×26" in xml


def test_wide_tables_switch_word_and_pdf_to_landscape(export_dir):
    md = _big_mto(batches=1, rows_per_batch=5)
    from docx import Document
    from docx.enum.section import WD_ORIENT
    doc = Document(E.export_docx(md, META))
    assert doc.sections[0].orientation == WD_ORIENT.LANDSCAPE
    assert b"/MediaBox [ 0 0 792 612 ]" in Path(E.export_pdf(md, META)).read_bytes()  # 11 x 8.5 in


def test_excel_numbers_are_numeric_and_formulas_stay_text(export_dir):
    md = "Table: Totals\n| Mark | Qty | Weight (kg) | Note |\n|---|---|---|---|\n| B1 | 12 | 1,234.5 | =SUM(A1) |\n"
    from openpyxl import load_workbook
    wb = load_workbook(E.export_xlsx(md, META))
    ws = wb["Totals"]
    assert ws["B2"].value == 12 and ws["C2"].value == 1234.5
    assert ws["D2"].value == "=SUM(A1)" and ws["D2"].data_type == "s"


def test_one_failing_format_does_not_block_the_others(export_dir, monkeypatch):
    def boom(content, meta):
        raise RuntimeError("broken")
    monkeypatch.setitem(E.EXPORTERS, "docx", boom)
    res = {r["format"]: r for r in E.generate_all_exports("# Hi\nBody", META)}
    assert res["docx"]["path"] == "" and "broken" in res["docx"]["error"]
    assert all(res[f]["path"] for f in ("markdown", "csv", "xlsx", "pdf"))


def test_arithmetic_asterisks_survive_but_emphasis_is_stripped():
    assert E._strip_md("Area = 2*3*4 = 24 ft2") == "Area = 2*3*4 = 24 ft2"
    assert E._strip_md("12 * 26 * 3.5 = 1,092 lb") == "12 * 26 * 3.5 = 1,092 lb"
    assert E._strip_md("**Weight:** *approx* `W12x26`") == "Weight: approx W12x26"
    assert E._xl_value("1*2.5*3") == ("1*2.5*3", None)


def test_lists_keep_their_numbers_and_nesting(export_dir):
    md = "Steps to verify:\n1. **Connection design** by EOR\n   - Moment frames\n2. Embed plates\n7. Galvanizing\n#### Deep heading"
    blocks = E._split_blocks(md)
    assert blocks[0] == ("p", "Steps to verify:") and blocks[-1] == ("h3", "Deep heading")
    items = [it for kind, text in blocks if kind == "list" for it in E._list_items(text)]
    assert items == [(0, "1.", "**Connection design** by EOR"), (1, "•", "Moment frames"),
                     (0, "2.", "Embed plates"), (0, "7.", "Galvanizing")]
    import pymupdf
    text = pymupdf.open(E.export_pdf(md, META))[0].get_text()
    assert "7." in text and "Galvanizing" in text


def test_escaped_pipes_stay_in_their_cell():
    rows = E._table_rows("| Mark | Flag | Sheet |\n| B12 | RECONCILE \\| VIF | S-201 |")
    assert rows[1] == ["B12", "RECONCILE | VIF", "S-201"]


def test_pdf_rows_taller_than_a_page_split_instead_of_failing(export_dir):
    long = "word " * 300
    md = "| # | Question | Impact | Resolution | Priority | Sheet | Status |\n|---|---|---|---|---|---|---|\n" + \
         f"| 1 | {long} | {long} | {long} | HIGH | S-201 | OPEN |\n"
    import pymupdf
    assert pymupdf.open(E.export_pdf(md, META)).page_count >= 2


def test_csv_neutralises_formula_cells(export_dir):
    path = E.export_csv("| =A1+A2 | x |\n|---|---|\n| +12'-6\" | y |\n\n@ grid C", META)
    body = Path(path).read_text(encoding="utf-8-sig")
    assert "'=A1+A2 | x" in body and "'+12'-6\"\" | y" in body and "'@ grid C" in body


def test_excel_numbers_display_as_written(export_dir):
    md = "Table: Nums\n| A | B | C | D |\n|---|---|---|---|\n| 3.10 | 1,234.50 | 12345678901234567891 | 0.50 |\n"
    from openpyxl import load_workbook
    ws = load_workbook(E.export_xlsx(md, META))["Nums"]
    assert ws["A2"].value == 3.1 and ws["A2"].number_format == "0.00"
    assert ws["B2"].value == 1234.5 and ws["B2"].number_format == "#,##0.00"
    assert ws["C2"].value == "12345678901234567891"
    assert ws["D2"].number_format == "0.00"
